#!/usr/bin/env python3
"""Tracker Service — the tracker is the applications table; this exports it
and records outcomes.

  python -m pipeline.tracker export               write job_tracker.xlsx
  python -m pipeline.tracker set <job_id> interview --notes "call with CTO"
  python -m pipeline.tracker weekly               print the weekly rollup
  python -m pipeline.tracker show                 recent rows

Success is response rate per application sent, not applications sent — so the
weekly rollup puts applied, reviewed and skipped side by side per fit-score
bucket, with the response and interview rates for what was actually sent. That
is what calibrates the thresholds over time.
"""
from __future__ import annotations

import argparse
import os
import sys
from datetime import date

from . import config as pc
from .db import DB, STATUSES

COLUMNS = ["date", "company", "role", "source", "jd_url", "language", "account",
           "cv_variant", "cv_filename", "fit_score", "decision", "status", "notes",
           "job_id"]
SENT = ("applied", "rejected", "interview", "offer")
RESPONDED = ("rejected", "interview", "offer")


def bucket(fit) -> str:
    if fit is None:
        return "pre-screen"
    if fit >= 70:
        return ">=70"
    if fit >= 40:
        return "40-69"
    return "<40"


def weekly(db: DB) -> list[dict]:
    rows = db.q("SELECT date, fit_score, decision, status FROM applications")
    agg: dict[tuple[str, str], dict] = {}
    for r in rows:
        try:
            y, w, _ = date.fromisoformat(r["date"]).isocalendar()
        except (TypeError, ValueError):
            continue
        key = (f"{y}-W{w:02d}", bucket(r["fit_score"]))
        a = agg.setdefault(key, {"week": key[0], "bucket": key[1], "decided": 0,
                                 "skipped": 0, "reviewed": 0, "applied": 0,
                                 "responses": 0, "interviews": 0})
        a["decided"] += 1
        st = r["status"]
        if st == "skipped":
            a["skipped"] += 1
        if r["decision"] == "review":
            a["reviewed"] += 1
        if st in SENT:
            a["applied"] += 1
        if st in RESPONDED:
            a["responses"] += 1
        if st in ("interview", "offer"):
            a["interviews"] += 1
    out = []
    order = {">=70": 0, "40-69": 1, "<40": 2, "pre-screen": 3}
    for (_, _), a in sorted(agg.items(), key=lambda kv: (kv[0][0], order[kv[0][1]]),
                            reverse=False):
        a["response_rate"] = round(a["responses"] / a["applied"], 3) if a["applied"] else None
        a["interview_rate"] = round(a["interviews"] / a["applied"], 3) if a["applied"] else None
        out.append(a)
    return out


def funnel(db: DB) -> dict:
    st = dict(db.q("SELECT stage, COUNT(*) FROM jobs GROUP BY stage"))
    dec = dict(db.q("SELECT decision, COUNT(*) FROM applications GROUP BY decision"))
    return {
        "sourced": sum(st.values()),
        "filtered out": st.get("filtered", 0),
        "candidates": sum(v for k, v in st.items() if k not in ("sourced", "filtered")),
        "evaluated by model": db.one("SELECT COUNT(*) FROM matches")[0],
        "auto": dec.get("auto", 0),
        "review": dec.get("review", 0),
        "skip": dec.get("skip", 0),
        "applied": db.one(f"SELECT COUNT(*) FROM applications WHERE status IN "
                          f"({','.join('?' * len(SENT))})", *SENT)[0],
    }


def export(db: DB, path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True, color="FFFFFF")
    head = PatternFill("solid", fgColor="3D5A3D")
    tint = {"applied": "E8F2E8", "interview": "CFE8CF", "offer": "B5DDB5",
            "staged": "FFF4DC", "pending": "FFF4DC", "review": "FDEBD3",
            "rejected": "F6E0E0", "skipped": "F2F2F2"}

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # -- Applications: the ApplicationRecord, one row per processed job --------
    ws = wb.active
    ws.title = "Applications"
    rows = db.q(f"SELECT {', '.join(COLUMNS)} FROM applications "
                f"ORDER BY date DESC, fit_score DESC")
    sheet(ws, COLUMNS, [list(r) for r in rows],
          [11, 18, 38, 11, 30, 8, 8, 26, 30, 8, 9, 10, 50, 22])
    for i, r in enumerate(rows, start=2):
        url = r["jd_url"]
        if url:
            c = ws.cell(row=i, column=5)
            c.hyperlink, c.style = url, "Hyperlink"
        fill = tint.get(r["status"])
        if fill:
            ws.cell(row=i, column=12).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=i, column=13).alignment = Alignment(wrap_text=False)

    # -- Weekly rollup -----------------------------------------------------------
    wk = weekly(db)
    cols = ["week", "bucket", "decided", "skipped", "reviewed", "applied",
            "responses", "interviews", "response_rate", "interview_rate"]
    sheet(wb.create_sheet("Weekly"), cols, [[a[c] for c in cols] for a in wk],
          [10, 11, 9, 9, 9, 9, 10, 10, 13, 13])

    # -- Review queue --------------------------------------------------------------
    rq = db.q("SELECT a.fit_score, a.company, a.role, a.cv_variant, a.notes, "
              "a.jd_url, a.job_id FROM applications a WHERE a.status='review' "
              "ORDER BY a.fit_score DESC")
    sheet(wb.create_sheet("Review queue"),
          ["fit", "company", "role", "cv_variant", "why", "jd_url", "job_id"],
          [list(r) for r in rq], [6, 18, 38, 26, 60, 30, 22])

    # -- Funnel --------------------------------------------------------------------
    fn = funnel(db)
    sheet(wb.create_sheet("Funnel"), ["stage", "jobs"],
          [[k, v] for k, v in fn.items()], [22, 10])

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(f".{os.getpid()}.tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)          # never leave a half-written spreadsheet


def set_status(db: DB, job_id: str, status: str, notes: str = "") -> None:
    if status not in STATUSES:
        raise ValueError(f"status must be one of: {', '.join(STATUSES)}")
    row = db.one("SELECT notes FROM applications WHERE job_id=?", job_id)
    if not row:
        raise KeyError(f"no tracker row for {job_id}")
    new_notes = row["notes"] or ""
    if notes:
        new_notes = f"{new_notes} | {date.today()}: {notes}".strip(" |")
    with db.tx():
        db.upsert_application(job_id, status=status, notes=new_notes)
        db.event("status", job_id, status=status, notes=notes)


def main() -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("export")
    sub.add_parser("weekly")
    sh = sub.add_parser("show")
    sh.add_argument("-n", type=int, default=20)
    st = sub.add_parser("set")
    st.add_argument("job_id")
    st.add_argument("status", choices=STATUSES)
    st.add_argument("--notes", default="")
    args = ap.parse_args()

    pcfg = pc.load()
    db = DB(pcfg.db)
    if args.cmd == "export":
        export(db, pcfg.tracker_xlsx)
        print(f"wrote {pcfg.tracker_xlsx}")
    elif args.cmd == "weekly":
        for a in weekly(db):
            rr = "-" if a["response_rate"] is None else f"{a['response_rate']:.0%}"
            print(f"  {a['week']}  {a['bucket']:10s}  decided {a['decided']:3d}  "
                  f"skipped {a['skipped']:3d}  reviewed {a['reviewed']:3d}  "
                  f"applied {a['applied']:3d}  response {rr}")
    elif args.cmd == "show":
        for r in db.q("SELECT * FROM applications ORDER BY updated_at DESC LIMIT ?", args.n):
            fit = "  -" if r["fit_score"] is None else f"{r['fit_score']:3d}"
            print(f"  {r['status']:9s} {fit}  {r['company'][:16]:16s} "
                  f"{r['role'][:44]:44s} {r['job_id']}")
    elif args.cmd == "set":
        set_status(db, args.job_id, args.status, args.notes)
        export(db, pcfg.tracker_xlsx)
        print(f"{args.job_id} -> {args.status}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
