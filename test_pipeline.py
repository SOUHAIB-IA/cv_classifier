#!/usr/bin/env python3
"""Tests for the job pipeline. Self-contained: temp database, no network, the
model stubbed. Chrome-dependent checks are skipped when Chrome is absent.

Run:  ./.venv/bin/python test_pipeline.py
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import httpx

import cvrouter as cr
import matcher
import testkit
from pipeline import calibrate as CAL
from pipeline import config as pc
from pipeline import orchestrate as O
from pipeline import sources as S
from pipeline import submit as SUB
from pipeline import tailor as T
from pipeline import tracker as TR
from pipeline.db import DB, dedupe_key
from pipeline.prefilter import Screener
from pipeline.textutil import html_to_text

HERE = Path(__file__).resolve().parent
ok = True


def check(label, cond, detail=""):
    global ok
    print(("  PASS  " if cond else "  FAIL  ") + label + (f"   ({detail})" if detail and not cond else ""))
    ok = ok and bool(cond)


def setup():
    cfg, tmp = testkit.temp_config()
    pcfg = pc.load(HERE / "pipeline.example.toml")
    pcfg.db = tmp / "p.db"
    pcfg.lock_file = tmp / "p.lock"
    pcfg.applications_dir = tmp / "applications"
    pcfg.tracker_xlsx = tmp / "tracker.xlsx"
    pcfg.structured_dir = tmp / "structured"
    # a tiny CV collection so the screener has a skill vocabulary to read
    idx = cr.Index(cfg)
    recs = {}
    for i, skills in enumerate([["Python", "PyTorch", "LangChain", "RAG", "Docker"],
                                ["Python", "SQL", "Docker", "FastAPI", "MLflow"],
                                ["Python", "TensorFlow", "Kubernetes", "LLM"]]):
        rel = f"2-Graduate/AI-ML-Engineering/cv-{i}.pdf"
        testkit.make_pdf(cfg.cv_root / rel, [f"Ada {i}", "AI Engineer"] + testkit.SAMPLE_TEXT[3:])
        recs[rel] = cr.CVRecord(path=rel, branch="2-Graduate", specialty="AI-ML-Engineering",
                                role="AI-ML-Engineer", lang="EN", headline="AI Engineer",
                                summary="AI engineer", skills=skills, sig=f"s{i}")
    idx.commit(upserts=recs)
    return cfg, pcfg, tmp


def job(i, company="Acme", title="Machine Learning Engineer", location="Paris, France",
        jd=None, lang="en", source="ashby"):
    return {"id": f"{source}:{i}", "source": source, "board": company.lower(),
            "company": company, "title": title, "location": location,
            "jd_text": jd or "We build LLM and RAG systems in Python with PyTorch and Docker. "
                             "You will ship machine learning to production.",
            "apply_url": f"https://jobs.ashbyhq.com/{company.lower()}/{i}/application",
            "jd_url": f"https://jobs.ashbyhq.com/{company.lower()}/{i}",
            "language": lang, "posted_date": "2026-09-10"}


def fake_match(fit, variant="2-Graduate/AI-ML-Engineering/cv-0.pdf"):
    def fn(cfg, jd, **kw):
        return {"best_variant": variant, "fit_score": fit, "ats_score": fit,
                "fit_reason": "stub", "job_language": "en",
                "suggested_edits": [], "key_changes": [], "best": {"path": variant}}
    return fn


def main():
    # ------------------------------------------------------------ text/sources
    print("\n1. job-board parsing")
    gh = "&lt;div&gt;&lt;h2&gt;About&lt;/h2&gt;&lt;ul&gt;&lt;li&gt;Python&lt;/li&gt;&lt;/ul&gt;&lt;/div&gt;"
    t = html_to_text(gh)
    check("Greenhouse double-encoded HTML becomes text", "About" in t and "- Python" in t
          and "<" not in t, t)

    def handler(req: httpx.Request):
        if "lever" in req.url.host:
            return httpx.Response(200, json=[{
                "id": "L1", "text": "ML Engineer", "descriptionPlain": "Intro",
                "lists": [{"text": "Requirements", "content": "<li>PyTorch</li>"}],
                "additionalPlain": "Benefits", "categories": {"location": "Paris"},
                "workplaceType": "hybrid", "hostedUrl": "https://jobs.lever.co/x/L1",
                "applyUrl": "https://jobs.lever.co/x/L1/apply", "createdAt": 1757500000000}])
        return httpx.Response(200, json={"jobs": [{
            "id": "A1", "title": "AI Engineer", "location": "Paris", "isRemote": True,
            "descriptionPlain": "Build agents", "jobUrl": "u", "applyUrl": "u/application",
            "publishedAt": "2026-09-01T10:00:00+00:00", "isListed": True},
            {"id": "A2", "title": "Hidden", "isListed": False}]})

    with httpx.Client(transport=httpx.MockTransport(handler)) as c:
        lv = S.lever(c, "x")[0]
        ab = S.ashby(c, "y")
    check("Lever joins intro, requirement lists and closing",
          all(x in lv["jd_text"] for x in ("Intro", "Requirements", "PyTorch", "Benefits")))
    check("Lever keeps the work mode with the place", lv["location"] == "Paris (hybrid)", lv["location"])
    check("Lever epoch-millis date parsed", lv["posted_date"].startswith("2025-09"), lv["posted_date"])
    check("Ashby unlisted postings dropped", len(ab) == 1)
    check("Ashby remote flag kept", "remote" in ab[0]["location"])

    # -------------------------------------------------------------- prefilter
    cfg, pcfg, tmp = setup()
    scr = Screener(pcfg, cfg)
    print("\n2. pre-screen")
    for title, want in [("Marketing Generalist - Saudi Arabia", False),
                        ("Internal AI Tools Engineer", True), ("Machine Learning Engineer", True),
                        ("Senior ML Engineer", False), ("Stage Data Science", False),
                        ("Ingénieur IA", True)]:
        check(f"title '{title}' -> {'kept' if want else 'dropped'}",
              scr._title_ok(title)[0] == want)
    for loc, want in [("Middle East, Dubai, Riyadh (remote)", False),
                      ("France, Europe, Paris (remote)", True), ("Remote", True),
                      ("Remote - US", False), ("San Francisco, CA | Paris", True)]:
        check(f"location '{loc}' -> {'kept' if want else 'dropped'}",
              scr._location_ok(loc)[0] == want)
    hi, _ = scr.score(job(1))
    lo, _ = scr.score(job(2, title="Software Engineer", jd="Ruby on Rails and PHP."))
    check("score follows the skills in YOUR CVs", hi > lo, f"{hi} vs {lo}")

    # ------------------------------------------------------------ orchestrate
    print("\n3. routing")
    db = DB(pcfg.db)
    with db.tx():
        for i, fit_title in enumerate(["ML Engineer A", "ML Engineer B", "ML Engineer C"]):
            db.upsert_job(job(i, company=f"Co{i}", title=fit_title))
    fits = iter([85, 55, 20])
    O.route(db, pcfg, cfg, match_fn=lambda *a, **k: fake_match(next(fits))(*a, **k),
            max_matches=5, log=lambda *a: None)
    got = {r["company"]: (r["decision"], r["status"]) for r in db.q("SELECT * FROM applications")}
    check("fit 85 -> auto / pending", got.get("Co0") == ("auto", "pending"), got.get("Co0"))
    check("fit 55 -> review", got.get("Co1") == ("review", "review"), got.get("Co1"))
    check("fit 20 -> skip", got.get("Co2") == ("skip", "skipped"), got.get("Co2"))

    with db.tx():
        db.upsert_job(job(9, company="Co0", title="ML Engineer A (H/F)"))
    calls = {"n": 0}

    def counted(*a, **k):
        calls["n"] += 1
        return fake_match(90)(*a, **k)

    O.route(db, pcfg, cfg, match_fn=counted, max_matches=5, log=lambda *a: None)
    check("duplicate guard: same company+title never evaluated again", calls["n"] == 0)
    check("dedupe key ignores (H/F) and case",
          dedupe_key("Co0", "ML Engineer A (H/F)") == dedupe_key("co0", "ml engineer a"))

    print("\n4. regional twins")
    with db.tx():
        for i, loc in enumerate(["Dubai, UAE (remote)", "France, Paris (remote)",
                                 "Singapore (remote)"]):
            db.upsert_job(job(20 + i, company="Cohere", title="Forward Deployed Engineer",
                              location=loc))
    seen = []
    O.route(db, pcfg, cfg, max_matches=5, log=lambda *a: None,
            match_fn=lambda c, jd, **k: (seen.append(k["location"]), fake_match(60)(c, jd, **k))[1])
    check("one role posted three times is evaluated once", len(seen) == 1, seen)
    check("the posting kept is the best-placed one", seen and "France" in seen[0], seen)

    print("\n5. budget and crash safety")
    pcfg.daily_match_budget = db.matches_today() + 1
    with db.tx():
        for i in range(3):
            db.upsert_job(job(40 + i, company=f"Budget{i}", title="ML Engineer"))
    calls["n"] = 0
    O.route(db, pcfg, cfg, match_fn=counted, max_matches=5, log=lambda *a: None)
    check("daily budget caps evaluations", calls["n"] == 1, calls["n"])

    pcfg.daily_match_budget = 1000
    with db.tx():
        db.upsert_job(job(60, company="Crashy", title="ML Engineer"))

    def boom(*a, **k):
        raise ValueError("model returned garbage")

    for _ in range(3):
        O.route(db, pcfg, cfg, match_fn=boom, max_matches=10, log=lambda *a: None)
    st = db.one("SELECT stage FROM jobs WHERE id='ashby:60'")["stage"]
    charged = db.one("SELECT COUNT(*) FROM events WHERE kind='match' AND job_id='ashby:60'")[0]
    check("a crashing evaluation is parked after 2 attempts, not retried forever",
          st == "error" and charged == 2, f"stage={st} charged={charged}")

    with db.tx():
        db.upsert_job(job(70, company="Dry", title="ML Engineer"))
    before = db.one("SELECT COUNT(*) FROM events")[0]
    O.route(db, pcfg, cfg, dry=True, match_fn=counted, max_matches=5, log=lambda *a: None)
    check("dry run writes nothing", db.one("SELECT COUNT(*) FROM events")[0] == before)

    # ------------------------------------------------------------------ tailor
    print("\n6. tailoring")
    base = {"name": "Ada Lovelace", "headline": "AI Engineer",
            "summary": "Engineer integrating LLMs with external APIs.",
            "contact": {}, "sections": [
                {"title": "Experience", "kind": "experience", "items": [
                    {"heading": "AI Engineer", "org": "Acme", "dates": "2025",
                     "bullets": ["Built a RAG pipeline with LangChain."]}]},
                {"title": "Skills", "kind": "skills", "groups": [
                    {"label": "Pipeline tools", "items": ["Airflow"]},
                    {"label": "Cloud", "items": ["Docker", "Azure"]}]}]}
    edits = [
        {"section": "Profile", "current": "integrating LLMs with external APIs",
         "suggested": "integrating LLMs through REST APIs"},
        {"section": "Skills — Cloud", "current": "Docker | Azure", "suggested": "Docker | REST API | Azure"},
        {"section": "Experience", "current": "Led a team of 40 engineers", "suggested": "CTO of 40"},
        {"section": "Profile", "current": "(absent)", "suggested": "Blockchain expert"},
        {"section": "Experience", "current": "Built a RAG pipeline",
         "suggested": "Built a RAG pipeline " + "x" * 400},
    ]
    out, rep = T.apply_edits(base, edits)
    st = [r["status"] for r in rep]
    check("anchored rewording applied", "through REST APIs" in out["summary"])
    check("skills group with a 'line'-like label handled as skills",
          out["sections"][1]["groups"][1]["items"] == ["Docker", "REST API", "Azure"])
    check("invented experience refused", st[2] == "skipped")
    check("blind addition to the profile refused", st[3] == "skipped")
    check("a rewrite disguised as an edit refused", st[4] == "skipped")
    check("the cached base CV is never mutated", "external APIs" in base["summary"])
    _, prot = T.apply_edits(base, [{"section": "Experience", "current": "Acme 2025",
                                    "suggested": "Google 2020"}])
    check("dates and employers are protected, and said so",
          prot[0]["status"] == "skipped" and "protected" in prot[0]["reason"], prot[0]["reason"])

    reordered = T.reorder(json.loads(json.dumps(base)), "graduate")
    check("graduate CVs lead with experience", reordered["sections"][0]["kind"] == "experience")

    pdf_text = "Ada Lovelace AI Engineer Engineer integrating LLMs with external APIs " \
               "AI Engineer Acme 2025 Built a RAG pipeline with LangChain Airflow Docker Azure"
    check("full extraction passes the coverage check", T.coverage(pdf_text, base) >= 0.85)
    check("lossy extraction fails it", T.coverage(pdf_text, {"name": "Ada Lovelace"}) < 0.85)

    try:
        chrome = T.find_chrome(pcfg)
    except T.TailorError:
        chrome = None
    if chrome:
        big = json.loads(json.dumps(base))
        big["sections"][0]["items"] = [
            {"heading": f"Role {i}", "org": "Org", "dates": "2020", "location": "Paris",
             "bullets": [f"Delivered outcome number {j} for project {i} with measurable impact."
                         for j in range(5)]} for i in range(7)]
        pdf = tmp / "fit.pdf"
        pages, step = T.render_fitted(big, "en", pdf, chrome, max_pages=1)
        check("a dense CV is tightened onto one page", pages == 1, f"{pages} pages")
        rep_ok = [{"status": "applied", "suggested": "Delivered outcome number 3"}]
        check("the PDF is read back and verified", T.verify(pdf, big, rep_ok, max_pages=1) == [])
        rep_bad = [{"status": "applied", "suggested": "Quantum gravity specialist"}]
        check("verification catches an edit missing from the PDF",
              T.verify(pdf, big, rep_bad, max_pages=1) != [])
    else:
        print("  skip  render checks (no Chrome)")

    # ------------------------------------------------------------------ submit
    print("\n7. submission guard rails")
    src = (HERE / "pipeline" / "submit.py").read_text()
    code = "\n".join(l for l in src.splitlines() if not l.lstrip().startswith(("#", '"', "-")))
    check("the submission assistant never calls click()", ".click(" not in code
          and ".dblclick(" not in code and ".tap(" not in code)
    check("it never presses Enter (which submits forms)", "press(" not in code)
    for url, want in [("https://jobs.lever.co/x/1/apply", 1),
                      ("https://job-boards.greenhouse.io/x/jobs/1", 1),
                      ("https://jobs.ashbyhq.com/x/1/application", 1),
                      ("https://www.linkedin.com/jobs/view/1", 2),
                      ("https://fr.indeed.com/viewjob?jk=1", 2),
                      ("https://evil.jobs.lever.co.attacker.com/x", 2)]:
        check(f"tier {want}: {url[:44]}", SUB.tier(url, pcfg.submit_hosts) == want)

    # ------------------------------------------------------------------ tracker
    print("\n8. tracker")
    TR.export(db, pcfg.tracker_xlsx)
    from openpyxl import load_workbook
    wb = load_workbook(pcfg.tracker_xlsx)
    check("xlsx has the four sheets", wb.sheetnames ==
          ["Applications", "Weekly", "Review queue", "Funnel"], wb.sheetnames)
    check("xlsx columns follow the ApplicationRecord",
          [c.value for c in wb["Applications"][1]][:13] == TR.COLUMNS[:13])
    try:
        TR.set_status(db, "ashby:0", "hired-maybe")
        check("unknown status rejected", False)
    except ValueError:
        check("unknown status rejected", True)
    TR.set_status(db, "ashby:0", "interview", "call with CTO")
    wk = [w for w in TR.weekly(db) if w["bucket"] == ">=70"]
    check("weekly rollup counts the interview in its fit bucket",
          wk and wk[-1]["interviews"] == 1 and wk[-1]["response_rate"] == 1.0, wk)

    # --------------------------------------------------------------- calibrate
    print("\n9. calibration")
    pairs = [(88, "apply"), (82, "apply"), (76, "apply"), (66, "maybe"), (58, "maybe"),
             (51, "maybe"), (35, "no"), (22, "no"), (73, "maybe"), (90, "apply")]
    agree, _, a, r = CAL.best_thresholds(pairs)[0]
    check("finds thresholds that reproduce the labels", agree >= 0.9, (agree, a, r))

    # ----------------------------------------------------------------- matcher
    print("\n10. matcher")
    check("French ad detected", matcher.guess_language(
        "Nous recherchons un ingénieur pour notre équipe, avec une expérience en IA.") == "fr")
    check("English ad detected", matcher.guess_language(
        "We are looking for an engineer to join our team with experience in AI.") == "en")
    real = cr.ask_json
    try:
        # Tell the two calls apart by what they are sent: only the shortlist
        # call carries the CV library (both system prompts say "shortlist").
        cr.ask_json = lambda c, s, u, **k: ({"shortlist": ["nope.pdf"]} if "=== CV LIBRARY ===" in u
                                            else {"best": {"path": "made-up.pdf"}, "fit_score": "71",
                                                  "ats_score": 140, "key_changes": [{"x": 1}]})
        idx = cr.Index(cfg)
        res = matcher.match_job(cfg, "We need Python.", idx=idx, text_of=lambda rel: "cv text")
        check("a CV the model invents is replaced by a real one", res["best_variant"] in idx.records)
        check("scores are clamped to 0-100 integers",
              res["fit_score"] == 71 and res["ats_score"] == 100)
        check("malformed edits are dropped", res["suggested_edits"] == [])
    finally:
        cr.ask_json = real

    # ------------------------------------------------------ photo and em dash
    print("\n12. photo, and never an em dash")
    check("the CV template has no em dash", "—" not in (HERE / "templates" / "cv.html").read_text())
    check("the model is told never to write one", "Never use the em dash" in matcher.ANALYSE_SYSTEM)
    for raw, want in [("Engineer — Acme", "Engineer, Acme"), ("LLM—RAG", "LLM-RAG"),
                      ("Essentials – Coursera", "Essentials, Coursera"),
                      ("09/2023 – 06/2026", "09/2023 – 06/2026")]:
        check(f"'{raw}' -> '{want}'", T.no_em_dash(raw) == want, T.no_em_dash(raw))
    dashy = {"name": "Ada", "headline": "AI Engineer — LLM", "summary": "x",
             "sections": [{"title": "Exp", "kind": "experience", "items": [
                 {"heading": "Engineer", "org": "Acme", "bullets": ["Built it — fast"]}]}]}
    html = T.render_html(dashy, "en", photo="data:image/png;base64,AAAA")
    body = html.split("</head>", 1)[1]
    check("nothing rendered contains an em dash, whatever the text says", "—" not in body)
    check("the photo is rendered when there is one", 'class="photo"' in html)
    check("…and left out when the CV turns it off",
          'class="photo"' not in T.render_html({**dashy, "show_photo": False}, "en",
                                              photo="data:image/png;base64,AAAA"))
    check("…and when there is no photo", 'class="photo"' not in T.render_html(dashy, "en"))
    pcfg_ph = pc.load(HERE / "pipeline.example.toml")
    pcfg_ph.photo = tmp / "me.png"
    (tmp / "me.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 32)
    pcfg_ph.photo_langs = ["fr"]
    check("photo limited to the languages chosen",
          T.photo_uri(pcfg_ph, "fr") is not None and T.photo_uri(pcfg_ph, "en") is None)

    # ----------------------------------------------------- a CV that looks right
    print("\n13. typography and contact line")
    if chrome:
        import subprocess as sp
        pro = {"name": "Ada Lovelace", "headline": "AI Engineer",
               "contact": {"email": "ada@example.org", "phone": "+33 6 00 00 00 00",
                           "links": [{"label": "LinkedIn",
                                      "url": "https://www.linkedin.com/in/ada-lovelace-853a011a5"},
                                     {"label": "Portfolio", "url": "https://ada-lovelace.vercel.app/"},
                                     {"label": "GitHub", "url": "https://github.com/ADA-LOVELACE"}]},
               "summary": "Engineer.", "sections": [
                   {"title": "Experience", "kind": "experience", "items": [
                       {"heading": "AI Engineer", "org": "Acme", "dates": "2025", "bullets": ["Built it."]}]}]}
        fonts_want = {"calibri": "Carlito", "cambria": "Caladea", "arial": "LiberationSans"}
        installed = sp.run(["fc-list", ":", "family"], capture_output=True, text=True).stdout
        for style, face in fonts_want.items():
            if face.replace("LiberationSans", "Liberation Sans") not in installed:
                print(f"  skip  {style} (font not installed)")
                continue
            out = tmp / f"font-{style}.pdf"
            T.html_to_pdf(T.render_html({**pro, "style": style}, "en"), out, chrome)
            emb = sp.run(["pdffonts", str(out)], capture_output=True, text=True).stdout
            check(f"{style}: the real typeface is embedded, not a fallback",
                  face in emb and "Tinos" not in emb, emb.splitlines()[2:4])
        out = tmp / "contact.pdf"
        T.html_to_pdf(T.render_html(pro, "en"), out, chrome)
        txt = cr.pdf_text(out, layout=False)
        check("no markup leaks into the contact line", "<span" not in txt and "&" not in txt)
        check("addresses are read whole, never split on their hyphen",
              all(u in txt for u in ["linkedin.com/in/ada-lovelace-853a011a5",
                                     "ada-lovelace.vercel.app", "github.com/ADA-LOVELACE"]))
        check("links stay clickable in the PDF", len(T.pdf_links(out)) == 3)
    doc_l = {"contact": {"links": [{"label": "LinkedIn"}, {"label": "Protfolio"}]}}
    real_links = T.pdf_links
    T.pdf_links = lambda p: ["https://www.linkedin.com/in/x", "https://x.vercel.app/",
                             "https://github.com/X"]
    try:
        T.attach_links(doc_l, Path("/dev/null"))
    finally:
        T.pdf_links = real_links
    urls = [l.get("url") for l in doc_l["contact"]["links"]]
    check("real addresses are attached to their labels, even a misspelt one",
          urls == ["https://www.linkedin.com/in/x", "https://x.vercel.app/", "https://github.com/X"], urls)

    # ------------------------------------------------------------ web interface
    print("\n11. web interface: read, edit, validate")
    from fastapi.testclient import TestClient
    import portal
    import webui
    # point the UI at the test database — a connection per request, as in production
    webui._ctx = lambda: (pcfg, cfg, DB(pcfg.db))
    client = TestClient(portal.app)
    H = {"X-CV-Router": "1"}

    check("dashboard state answers", client.get("/api/pipeline/state").status_code == 200)
    r = client.post("/api/jobs", json={"company": "UI Co", "title": "ML Engineer",
                                       "jd": "We build LLM systems in Python. " * 5})
    check("a change without the X-CV-Router header is refused", r.status_code == 403)
    r = client.post("/api/jobs", headers=H, json={"company": "UI Co", "title": "ML Engineer",
                                                  "jd": "We build LLM systems in Python. " * 5})
    jid = r.json().get("job_id", "")
    check("a job can be added by hand", r.status_code == 200 and jid.startswith("manual:"))

    variant = "2-Graduate/AI-ML-Engineering/cv-0.pdf"
    lines = ["Ada 0", "AI Engineer"] + testkit.SAMPLE_TEXT[3:]
    extracted = {"name": "Ada 0", "headline": "AI Engineer", "contact": {},
                 "summary": " ".join(lines[2:]), "sections": [
                     {"title": "Skills", "kind": "skills",
                      "groups": [{"label": "Stack", "items": ["Python", "Docker"]}]}]}
    real_ask = cr.ask_json
    cr.ask_json = lambda *a, **k: json.loads(json.dumps(extracted))
    try:
        O.evaluate_job(db, pcfg, cfg, dict(db.one("SELECT * FROM jobs WHERE id=?", jid)),
                       match_fn=lambda c, jd, **k: {
                           "best_variant": variant, "fit_score": 88, "ats_score": 80,
                           "job_language": "en", "fit_reason": "stub",
                           "suggested_edits": [{"section": "Title", "current": "AI Engineer",
                                                "suggested": "AI Engineer (LLM)"}]})
        if chrome:
            r = client.post(f"/api/job/{jid}/tailor", headers=H)
            check("the CV is tailored from the UI", r.status_code == 200, r.text[:120])
            st = db.one("SELECT status FROM applications WHERE job_id=?", jid)["status"]
            check("a tailored CV starts as a draft you must read", st == "draft", st)

            data = client.get(f"/api/job/{jid}").json()
            check("the editor gets the document, the original and the edit report",
                  data["cv"]["doc"]["headline"] == "AI Engineer (LLM)"
                  and data["cv"]["base"]["headline"] == "AI Engineer"
                  and data["cv"]["meta"]["edits"][0]["status"] == "applied")

            doc = data["cv"]["doc"]
            doc["headline"] = "Machine Learning Engineer — edited by hand"
            html = client.post(f"/api/job/{jid}/preview", json={"doc": doc}).text
            check("preview renders the edited document", "edited by hand" in html)
            r = client.post(f"/api/job/{jid}/cv", headers=H, json={"doc": doc})
            check("saving regenerates the PDF", r.status_code == 200 and r.json()["pages"] == 1,
                  r.text[:120])
            pdf = client.get(f"/job/{jid}/cv.pdf")
            check("the PDF served is the edited one", pdf.status_code == 200)
            p = tmp / "served.pdf"
            p.write_bytes(pdf.content)
            check("…and an ATS reads the hand edit in it", "edited by hand" in cr.pdf_text(p))

            r = client.post(f"/api/job/{jid}/status", headers=H, json={"action": "validate"})
            check("validating moves the draft to ready-to-submit",
                  r.status_code == 200 and r.json()["status"] == "staged")
            r = client.post(f"/api/job/{jid}/status", headers=H, json={"action": "applied"})
            check("outcomes can be recorded from the UI", r.json()["status"] == "applied")
        else:
            print("  skip  UI tailoring checks (no Chrome)")
    finally:
        cr.ask_json = real_ask

    shutil.rmtree(tmp)
    print("\n" + ("ALL PASS" if ok else "FAILURES ABOVE"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
