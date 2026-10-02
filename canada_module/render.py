"""Rendering a Canadian resume: Letter PDF, and a .docx beside it.

The PDF goes through headless Chrome, the same mechanism the rest of the repo
uses, but through this module's own Letter template. The DOCX is built from the
same structured document rather than converted from the PDF, so its text is real
text in reading order — which is the only property an ATS cares about.
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

import browsers

HERE = Path(__file__).resolve().parent
TEMPLATES = HERE / "templates"

# Metric clones first: these are what a Linux box actually has, and they lay out
# identically to the names a Canadian career centre lists.
FONTS = {
    "calibri": '"Carlito", "Calibri", "Liberation Sans", Arial, sans-serif',
    "arial": '"Liberation Sans", Arial, Helvetica, sans-serif',
    "cambria": '"Caladea", "Cambria", Georgia, serif',
    "garamond": '"EB Garamond", Garamond, "Times New Roman", serif',
}
DEFAULT_FONT = "calibri"

# Loosest first. The floor is 10pt because UBC and UTM both put the body at
# 10-12pt: a Canadian resume that needs less is a resume with too much on it, and
# the answer is to cut content rather than to keep shrinking.
DENSITY = [
    {"fs": 11.0, "lh": 1.38, "margin": "19mm 18mm 17mm", "h2": "4.6mm",
     "item": "2.4mm", "h1": 19.0, "hl": 11.5, "li": ".6mm"},
    {"fs": 10.5, "lh": 1.32, "margin": "17mm 16mm 15mm", "h2": "4.0mm",
     "item": "2.0mm", "h1": 18.0, "hl": 11.0, "li": ".5mm"},
    {"fs": 10.0, "lh": 1.26, "margin": "15mm 14mm 13mm", "h2": "3.4mm",
     "item": "1.6mm", "h1": 17.0, "hl": 10.5, "li": ".4mm"},
]

LABELS = {"en": {"summary": "PROFESSIONAL SUMMARY"},
          "fr_qc": {"summary": "SOMMAIRE PROFESSIONNEL"}}


def env() -> Environment:
    return Environment(loader=FileSystemLoader(str(TEMPLATES)),
                       autoescape=select_autoescape(["html"]),
                       trim_blocks=True, lstrip_blocks=True)


def render_html(doc: dict, lang: str, density: dict | None = None,
                style: str = DEFAULT_FONT) -> str:
    return env().get_template("canada_cv.html").render(
        doc=doc, lang=lang, d=density or DENSITY[0],
        font=FONTS.get(style, FONTS[DEFAULT_FONT]),
        labels=LABELS["en" if lang == "en" else "fr_qc"])


def find_chrome(chrome_bin: str = "") -> str:
    p = browsers.find_chrome(chrome_bin)
    if not p:
        raise RuntimeError("no Chrome/Chromium found to render the PDF")
    return p


def html_to_pdf(html: str, out: Path, chrome: str) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "cv.html"
        src.write_text(html, encoding="utf-8")
        r = subprocess.run(
            [chrome, "--headless=new", "--disable-gpu", "--no-sandbox",
             "--no-pdf-header-footer", f"--user-data-dir={td}/profile",
             f"--print-to-pdf={out}", src.as_uri()],
            capture_output=True, text=True, timeout=120)
        if not out.is_file():
            raise RuntimeError(f"Chrome produced no PDF: {r.stderr[-400:]}")


def page_count(pdf: Path) -> int:
    out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True,
                         timeout=30).stdout
    for ln in out.splitlines():
        if ln.startswith("Pages:"):
            return int(ln.split(":", 1)[1].strip() or 0)
    return 0


def render_pdf(doc: dict, lang: str, out: Path, chrome: str,
               prefer_pages: int = 1, max_pages: int = 2,
               style: str = DEFAULT_FONT) -> tuple[int, int, str]:
    """Render at the loosest density that reaches `prefer_pages`.

    Two targets, because the Canadian rule is a range: one page is better and
    two are allowed. Tightening to 10pt to reach one page is worth it; going
    below 10pt is not, since UBC and UTM both floor the body there. So the
    search stops at the floor, keeps the loosest render that at least fits
    `max_pages`, and says plainly that the remaining fix is to cut content.

    Returns (pages, density step, note). An empty note means it fit as asked.
    """
    best: tuple[int, int, str] | None = None
    for step, d in enumerate(DENSITY):
        html_to_pdf(render_html(doc, lang, d, style), out, chrome)
        pages = page_count(out)
        if pages <= prefer_pages:
            return pages, step, ""
        if pages <= max_pages and best is None:
            best = (pages, step, "")
    if best is not None:
        # re-render the one we settled on, since the loop moved past it
        html_to_pdf(render_html(doc, lang, DENSITY[best[1]], style), out, chrome)
        return (best[0], best[1],
                f"{best[0]} pages, within the maximum of {max_pages} but not "
                f"the {prefer_pages} preferred. Cut the weakest content to "
                f"reach {prefer_pages}; the font floor is already at "
                f"{DENSITY[-1]['fs']}pt.")
    pages, step = page_count(out), len(DENSITY) - 1
    return (pages, step,
            f"{pages} pages at the {DENSITY[-1]['fs']}pt floor, over the "
            f"maximum of {max_pages}: cut content rather than shrink further")


# --------------------------------------------------------------------- DOCX
def render_docx(doc: dict, lang: str, out: Path) -> None:
    """The same document as Word, single column, no tables, no text boxes.

    Built from the structured document, not converted from the PDF: an ATS reads
    paragraphs in order, and that is exactly what this writes.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)   # US Letter
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(0.75))

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    def para(text="", *, bold=False, size=10.5, caps=False, space_before=0,
             space_after=2, colour=None):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text.upper() if caps else text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = "Calibri"
        if colour:
            r.font.color.rgb = RGBColor(*colour)
        return p

    para(doc.get("name", ""), bold=True, size=18, space_after=1)
    if doc.get("headline"):
        para(doc["headline"], bold=True, size=11, space_after=1)

    c = doc.get("contact", {}) or {}
    bits = [c.get("email"), c.get("phone"), c.get("location")]
    bits += [l.get("url") or l.get("label")
             for l in (c.get("links") or []) if l.get("url") or l.get("label")]
    if c.get("work_authorization"):
        bits.append(c["work_authorization"])
    para(" | ".join(b for b in bits if b), size=10, space_after=6)

    def heading(text):
        para(text, bold=True, size=10.5, caps=True, space_before=8,
             space_after=3, colour=(0x1F, 0x2A, 0x3A))

    if doc.get("summary"):
        heading(LABELS["en" if lang == "en" else "fr_qc"]["summary"])
        para(doc["summary"])

    for s in doc.get("sections", []):
        if s.get("hide"):
            continue
        heading(s.get("title", ""))
        if s.get("kind") == "skills":
            colon = "\u00a0: " if lang == "fr_qc" else ": "   # French spacing
            for g in s.get("groups") or []:
                label = f"{g['label']}{colon}" if g.get("label") else ""
                para(label + ", ".join(g.get("items") or []))
        elif s.get("kind") == "list":
            for line in s.get("lines") or []:
                para(str(line))
        else:
            for it in s.get("items") or []:
                if it.get("hide"):
                    continue
                meta = " · ".join(x for x in (it.get("location"),
                                                  it.get("dates")) if x)
                head = it.get("heading", "")
                if it.get("org"):
                    head += f", {it['org']}"
                if meta:
                    head += f" · {meta}"
                para(head, bold=True, space_after=1)
                for b in it.get("bullets") or []:
                    p = d.add_paragraph(str(b), style="List Bullet")
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(10.5)
                        r.font.name = "Calibri"

    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
