#!/usr/bin/env python3
"""Tests for the Canada converter.

The property that matters most is the last one: nothing the converter writes is
a fact the source did not already carry. Everything else — spelling, dates,
section order — is checked item by item, and then end to end by running the
linter over the converted document.

Run:  ./.venv/bin/python canada_module/tests/test_converter.py
"""
from __future__ import annotations

import dataclasses
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(ROOT))

from canada_module import config as ccfg          # noqa: E402
from canada_module import convert as CV           # noqa: E402
from canada_module import render as R             # noqa: E402
from canada_module.checker import check_doc, check_pdf  # noqa: E402
from canada_module.rules import load_rules        # noqa: E402

FIX = HERE / "fixtures"
ok = True


def check(label, cond, detail=""):
    global ok
    print(("  PASS  " if cond else "  FAIL  ") + label
          + (f"   ({detail})" if detail and not cond else ""))
    ok = ok and bool(cond)


def words(s: str) -> set[str]:
    return {w.lower() for w in re.findall(r"[A-Za-zÀ-ÿ']{3,}", s)}


def doc_words(doc: dict) -> set[str]:
    from canada_module.checker import _doc_text
    return words(_doc_text(doc))


def find_chrome():
    for c in ("google-chrome", "google-chrome-stable", "chromium",
              "chromium-browser"):
        if shutil.which(c):
            return shutil.which(c)
    return None


def main() -> int:
    rules = load_rules()
    cfg = ccfg.load()
    good = json.loads((FIX / "compliant.cv.json").read_text(encoding="utf-8"))
    bad = json.loads((FIX / "noncompliant.cv.json").read_text(encoding="utf-8"))

    # ------------------------------------------------------------- 1. spelling
    print("\n1. spelling")
    log = CV.Changelog()
    check("modeling -> modelling",
          CV.respell("Data modeling and SQL", rules, "en", log, "x")
          == "Data modelling and SQL")
    check("case of the first letter is preserved",
          CV.respell("Modeling & storage", rules, "en", log, "x")
          == "Modelling & storage",
          CV.respell("Modeling & storage", rules, "en", log, "x"))
    check("-ize and -yze forms are left alone",
          CV.respell("Optimized, containerized and analyzed the organization",
                     rules, "en", log, "x")
          == "Optimized, containerized and analyzed the organization")
    check("'program' is not turned into 'programme'",
          CV.respell("built the program", rules, "en", log, "x")
          == "built the program")
    # A domain is not prose. Rewriting it would break the link.
    url = "see https://data-center.example.com/colors and a@center.io"
    check("a URL and an email survive untouched",
          CV.respell(url, rules, "en", log, "x") == url,
          CV.respell(url, rules, "en", log, "x"))
    check("Quebec French: e-mail -> courriel",
          "courriel" in CV.respell("adresse e-mail", rules, "fr_qc", log, "x"))
    check("every substitution is logged",
          any(c.what == "spelling" for c in log.changes))

    # ---------------------------------------------------------------- 2. dates
    print("\n2. dates")
    log = CV.Changelog()
    check("MM/YYYY -> Mon YYYY",
          CV.reformat_dates("02/2026 – 08/2026", rules, "en", log, "x")
          == "Feb 2026 – Aug 2026",
          CV.reformat_dates("02/2026 – 08/2026", rules, "en", log, "x"))
    check("ISO -> Mon YYYY",
          CV.reformat_dates("2023-09-01 - 2026-06-30", rules, "en", log, "x")
          == "Sep 2023 – Jun 2026",
          CV.reformat_dates("2023-09-01 - 2026-06-30", rules, "en", log, "x"))
    check("French months in the Quebec variant",
          CV.reformat_dates("02/2026 – 08/2026", rules, "fr_qc", log, "x")
          == "févr. 2026 – août 2026",
          CV.reformat_dates("02/2026 – 08/2026", rules, "fr_qc", log, "x"))
    # Inventing a month for "2023" would be inventing a fact.
    check("a bare year is left alone",
          CV.reformat_dates("2020 – 2023", rules, "en", log, "x")
          == "2020 – 2023",
          CV.reformat_dates("2020 – 2023", rules, "en", log, "x"))
    check("'Present' is written in the CV's language",
          CV.reformat_dates("01/2026 - Present", rules, "fr_qc", log, "x")
          .endswith("à ce jour"),
          CV.reformat_dates("01/2026 - Present", rules, "fr_qc", log, "x"))

    # ------------------------------------------------------------- 3. locations
    print("\n3. locations")
    log = CV.Changelog()
    # The City, Province rule is about where the CANDIDATE is. A job in Rabat is
    # in Morocco, and relabelling it "Rabat, ON" would be a lie.
    check("an employer's country is translated, not relabelled",
          CV.normalise_place("Rabat, Maroc", "en", log, "x") == "Rabat, Morocco")
    check("...and no province is invented for it",
          "ON" not in CV.normalise_place("Rabat, Maroc", "en", log, "x")
          and "QC" not in CV.normalise_place("Rabat, Maroc", "en", log, "x"))
    check("English becomes French in the Quebec variant",
          CV.normalise_place("El Jadida, Morocco", "fr_qc", log, "x")
          == "El Jadida, Maroc")
    for mode, want in (("city_province", "ON"), ("city_country", "Morocco"),
                       ("omit", "")):
        c = dataclasses.replace(cfg, location_mode=mode, city="Casablanca",
                                country="Morocco", province="ON",
                                city_default="Toronto")
        got = CV.contact_location(c, "en")
        check(f"contact location, mode {mode}", want in got if want else got == "",
              got)
    c = dataclasses.replace(cfg, location_mode="city_country",
                            overseas_suffix="open to relocation to Canada")
    check("the overseas note is appended, not asserted as a Canadian address",
          "open to relocation" in CV.contact_location(c, "en", rules))
    # A French CV with an English contact line is how this first came out: the
    # function read the config's own strings and ignored `lang`.
    fr_loc = CV.contact_location(c, "fr_qc", rules)
    check("the French CV gets a French contact line",
          "Maroc" in fr_loc and "relocalisation" in fr_loc
          and "Morocco" not in fr_loc, fr_loc)

    # ------------------------------------------- 4. stripping what Canada bars
    print("\n4. what gets stripped")
    out, log = CV.convert(bad, rules, cfg, "en")
    check("the photo is gone", "photo" not in out)
    titles = [s.get("title", "") for s in out["sections"]]
    check("the References section is gone",
          not any("REFERENCE" in t.upper() for t in titles), titles)
    check("the personal-details section is gone once emptied",
          not any("PERSONAL" in t.upper() for t in titles), titles)
    removed = " ".join(c.where + c.before for c in log.changes
                       if c.what.startswith("removed"))
    for want in ("birth", "Nationality", "Insurance", "Gender"):
        check(f"removal of {want} is logged", want.lower() in removed.lower(),
              removed[:200])
    check("every removal carries a reason",
          all(c.why for c in log.changes if c.what.startswith("removed")))

    # --------------------------------------------------------- 5. the sections
    print("\n5. sections")
    out, log = CV.convert(good, rules, cfg, "en")
    titles = [s.get("title", "") for s in out["sections"]]
    want = [rules.headings("en")[k] for k in rules.order("en")
            if k != "summary" and rules.headings("en")[k] in titles]
    check("sections end up in the configured order", titles == want,
          f"{titles} != {want}")
    check("skills come before experience on the English CV",
          titles.index("TECHNICAL SKILLS") < titles.index("PROFESSIONAL EXPERIENCE"))
    fr_out, _ = CV.convert(good, rules, cfg, "fr_qc")
    fr_titles = [s.get("title", "") for s in fr_out["sections"]]
    check("the Quebec CV follows the OQLF order, formation before experience",
          fr_titles.index("FORMATION") < fr_titles.index("EXPÉRIENCE PROFESSIONNELLE"),
          fr_titles)
    combined = {"name": "X", "contact": {"email": "a@b.co", "phone": "+1 514 555 0199"},
                "sections": [{"title": "LANGUAGES & CERTIFICATIONS", "kind": "list",
                              "lines": ["Languages: Arabic, French, English",
                                        "Certifications: Coursera 2023"]}]}
    out2, log2 = CV.convert(combined, rules, cfg, "en")
    t2 = [s["title"] for s in out2["sections"]]
    check("a combined languages+certifications section is split in two",
          "LANGUAGES" in t2 and "CERTIFICATIONS" in t2, t2)

    # --------------------------------------------- 6. what it refuses to decide
    print("\n6. what it refuses to decide for you")
    out, log = CV.convert(good, rules, cfg, "en")
    check("an empty equivalence_text is reported, not filled in",
          any("equivalence_text is empty" in u for u in log.unresolved),
          log.unresolved)
    check("...and nothing was written on the degree",
          not any("equivalent" in str(i.get("heading", "")).lower()
                  for s in out["sections"] for i in s.get("items") or []))
    out, log = CV.convert(bad, rules, cfg, "en")
    check("a 'Responsible for' bullet is reported, not rewritten",
          any("Responsible for" in u for u in log.unresolved), log.unresolved)
    kept = [b for s in out["sections"] for i in s.get("items") or []
            for b in i.get("bullets") or []]
    check("...and the bullet is left exactly as written",
          any(b.startswith("Responsible for") for b in kept), kept[:2])
    check("a section over the bullet cap is reported",
          any("bullets and the cap" in u for u in log.unresolved))
    c = dataclasses.replace(cfg, auth_mode="omit")
    out, _ = CV.convert(good, rules, c, "en")
    check("work authorization stays off the resume by default",
          "work_authorization" not in out.get("contact", {}))
    c = dataclasses.replace(cfg, auth_mode="one_line", authorized_to_work=True)
    out, log = CV.convert(good, rules, c, "en")
    check("...and is written only when you confirmed it yourself",
          "eligible to work in Canada"
          in out["contact"].get("work_authorization", ""))

    # ------------------------------------------------------ 7. no invented fact
    print("\n7. nothing invented")
    out, log = CV.convert(good, rules, cfg, "en")
    allowed = set()
    for lang in ("en", "fr_qc"):
        allowed |= words(" ".join(rules.headings(lang).values()))
        allowed |= words(" ".join(rules.d["dates"]["months"][lang]))
        allowed |= words(rules.d["dates"]["present"][lang])
        allowed |= words(" ".join(
            str(v) for v in rules.lang(lang).get("substitute", {}).values()))
    allowed |= words(CV.contact_location(cfg, "en"))
    allowed |= words(" ".join(CV.COUNTRY_EN.values()))
    allowed |= words(rules.d["work_authorization"]["templates"]["one_line"]["en"])
    allowed |= {"canada", "letter", "lang", "paper", "built"}
    new = doc_words(out) - doc_words(good) - allowed
    check("the converted document introduces no new content word", not new, new)

    # --------------------------------------------------------- 8. end to end
    print("\n8. end to end: the converter fixes what the linter flags")
    before = check_doc(bad, rules, "en")
    check("the source document fails the linter", not before.ok)
    out, log = CV.convert(bad, rules, cfg, "en")
    after = check_doc(out, rules, "en")
    check("the converted document passes it", after.ok, after.text())
    # Spelling and openers survive as warnings or unresolved items by design:
    # a bullet with no result in it cannot be fixed without a new fact.
    check("the remaining items are warnings, never failures",
          not after.fails, [f.message for f in after.fails])
    # It cannot make a resume compliant that is missing something only you can
    # supply, and it must not invent it.
    thin = {"name": "X", "contact": {}, "summary": "A summary.", "sections": []}
    out3, _ = CV.convert(thin, rules, cfg, "en")
    rep3 = check_doc(out3, rules, "en")
    check("a missing email is reported, not invented",
          any("email" in f.message for f in rep3.fails)
          and not out3["contact"].get("email"),
          [f.message for f in rep3.findings])

    chrome = find_chrome()
    if not chrome:
        print("  skip  render checks (no Chrome)")
    else:
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            out, log = CV.convert(good, rules, cfg, "en")
            pdf = td / "x.pdf"
            pages, step, note = R.render_pdf(out, "en", pdf, chrome,
                                             prefer_pages=1, max_pages=2)
            check("it renders", pdf.is_file() and pages >= 1)
            rep = check_pdf(pdf, rules, "en")
            check("the rendered PDF passes the Canada rules", rep.ok, rep.text())
            check("paper is letter", "paper_size" not in {f.check for f in rep.findings})
            check("one column", "column_count" not in {f.check for f in rep.findings})
            check("no embedded image, so no photo",
                  "forbidden_content" not in {f.check for f in rep.findings},
                  [f.message for f in rep.findings])
            # The font stack carries double quotes; autoescaping them yields
            # invalid CSS and Chrome silently falls back to a serif. It did.
            check("the font stack survived autoescaping",
                  "font_not_allowed" not in {f.check for f in rep.findings},
                  [f.message for f in rep.findings if f.check == "font_not_allowed"])
            docx = td / "x.docx"
            R.render_docx(out, "en", docx)
            from docx import Document
            d = Document(str(docx))
            check("the DOCX is US Letter",
                  round(d.sections[0].page_width.inches, 2) == 8.5
                  and round(d.sections[0].page_height.inches, 2) == 11.0)
            check("the DOCX has no tables", len(d.tables) == 0)
            check("the DOCX leads with the name, not a header",
                  d.paragraphs[0].text.strip() == out["name"])
            txt = "\n".join(p.text for p in d.paragraphs)
            check("the DOCX carries the experience bullets",
                  "water-consumption anomaly detection" in txt)

    print("\n" + ("ALL PASS" if ok else "FAILURES ABOVE"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
